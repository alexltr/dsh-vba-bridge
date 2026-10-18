import json
from openpyxl import load_workbook
def norm(v):
    return getattr(v, "text", v)
a = load_workbook(r"F:\MyProjects\dsh\vba_bridge\fixtures\sample.xlsm", data_only=False)
b = load_workbook(r"F:\MyProjects\dsh\vba_bridge\work\perf.xlsm", data_only=False)
diffs = []
for name in a.sheetnames:
    wa, wb = a[name], b[name]
    if (wa.max_row, wa.max_column) != (wb.max_row, wb.max_column):
        diffs.append({"sheet": name, "dims": [wa.max_row, wa.max_column, wb.max_row, wb.max_column]})
        continue
    for row in wa.iter_rows():
        for cell in row:
            if norm(cell.value) != norm(wb[cell.coordinate].value):
                diffs.append({"sheet": name, "cell": cell.coordinate, "a": repr(norm(cell.value)), "b": repr(norm(wb[cell.coordinate].value))})
print(json.dumps({"sheets_equal": a.sheetnames == b.sheetnames, "sheet_count": len(a.sheetnames), "real_diffs": diffs[:10], "diff_count": len(diffs)}, ensure_ascii=False, indent=2))
